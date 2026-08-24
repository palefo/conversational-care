"""Turn an uploaded file into plain text.

Three formats, three small readers — deliberately no document-conversion
framework. Everything here is best-effort per page/paragraph: one unreadable
page in a fifty-page PDF should cost that page, not the whole upload.

Nothing in this module touches the database or the network.
"""
from __future__ import annotations

import logging
import os
import re

logger = logging.getLogger(__name__)

# Extensions we accept. Kept in one place: the model validator, the form, the
# drag-and-drop UI and this reader must agree or a file gets in that nothing
# can read.
SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx")

# Upload ceiling. Well below what the extractors can handle — this is about
# keeping one file from monopolising a worker thread for minutes.
MAX_UPLOAD_BYTES = 25 * 1024 * 1024


class ExtractionError(Exception):
    """The file could not be read at all (wrong format, corrupt, encrypted)."""


def extension_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def is_supported(filename: str) -> bool:
    return extension_of(filename) in SUPPORTED_EXTENSIONS


def _normalise(text: str) -> str:
    """Tidy extracted text without changing what it says.

    PDF and DOCX extraction both produce ragged whitespace — hard-wrapped
    lines, runs of blank lines, non-breaking spaces. Left alone that noise ends
    up inside chunks and inside the vectors.
    """
    text = text.replace(" ", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_txt(path: str) -> str:
    """Read a plain-text/Markdown file, guessing the encoding if it isn't UTF-8.

    Multilingual knowledge bases mean files that are not UTF-8 are normal, not
    exotic: a Spanish or Portuguese export from an old system is often cp1252,
    and decoding it as UTF-8 either fails or mangles every accent.
    """
    with open(path, "rb") as fp:
        raw = fp.read()
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    try:
        import chardet
        guess = chardet.detect(raw) or {}
        encoding = guess.get("encoding")
        if encoding:
            return raw.decode(encoding, errors="replace")
    except Exception:  # pragma: no cover - chardet is optional at runtime
        logger.debug("chardet unavailable; falling back to latin-1")
    return raw.decode("latin-1", errors="replace")


def _extract_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - pinned in requirements
        raise ExtractionError("PDF support is not installed (add `pypdf`).") from exc

    try:
        reader = PdfReader(path)
    except Exception as exc:
        raise ExtractionError(f"This PDF could not be opened: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        # An empty password unlocks most "protected" PDFs; a real one we can't help.
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("This PDF is password-protected.") from exc

    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            logger.warning("Could not extract page %d of %s", number, path)
            pages.append("")
    return "\n\n".join(p for p in pages if p.strip())


def _extract_docx(path: str) -> str:
    """Read a .docx, including its tables.

    Tables carry a lot of the content in care documents (schedules, dosages),
    and `python-docx` does not include them in `document.paragraphs`, so they
    have to be walked separately or they simply vanish from the knowledge base.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractionError(
            "Word support is not installed (add `python-docx`)."
        ) from exc

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ExtractionError(
            "This file could not be opened as a Word document. Note that the "
            "older .doc format is not supported — save it as .docx first."
        ) from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                # Pipe-separated, on one line. Two reasons: the splitter is far
                # less likely to cut a row in half, and a retrieved extract
                # reading "Donepezil | 10 mg | at night" is unambiguous where
                # space-separated cells would run together. (Tabs would not
                # survive — _normalise collapses them to spaces.)
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(path: str, filename: str | None = None) -> str:
    """Extract plain text from ``path``, dispatching on ``filename``'s extension.

    Raises ``ExtractionError`` when the format is unsupported or the file is
    unreadable. An empty return means the file opened fine but held no text —
    a scanned PDF, typically — which the caller reports differently, because
    the fix (run OCR / upload a text PDF) is a different one.
    """
    ext = extension_of(filename or path)
    if ext in (".txt", ".md"):
        text = _extract_txt(path)
    elif ext == ".pdf":
        text = _extract_pdf(path)
    elif ext == ".docx":
        text = _extract_docx(path)
    else:
        raise ExtractionError(
            f"Unsupported file type '{ext or '?'}'. "
            f"Accepted: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    return _normalise(text)
